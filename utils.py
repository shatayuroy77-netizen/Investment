import os
import streamlit as st
from google import genai
import finnhub
from dotenv import load_dotenv
from docx import Document
import io
import re

# Load secret keys (API Keys) from .env file
load_dotenv()

# ==============================
# API CONFIGURATION
# ==============================
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
FINNHUB_API_KEY = os.getenv("FINNHUB_API_KEY")

# Setup Gemini Client (Latest google-genai standard)
gemini_client = None
if GEMINI_API_KEY:
    gemini_client = genai.Client(api_key=GEMINI_API_KEY)

# Setup Finnhub Client
finnhub_client = None
if FINNHUB_API_KEY:
    finnhub_client = finnhub.Client(api_key=FINNHUB_API_KEY)

# ==============================
# FINNHUB API FUNCTIONS
# ==============================
@st.cache_data(ttl=1800)  # Cached for 30 minutes to prevent redundant API calls
def get_company_profile(ticker: str):
    """Fetches basic company profile and details from Finnhub."""
    if not finnhub_client:
        return None
    try:
        profile = finnhub_client.company_profile2(symbol=ticker.upper())
        return profile if profile else None
    except Exception as e:
        st.error(f"Error fetching profile from Finnhub: {e}")
        return None

@st.cache_data(ttl=300)
def get_stock_quote(ticker: str):
    """Fetches current stock price and quotes from Finnhub."""
    if not finnhub_client:
        return None
    try:
        quote = finnhub_client.quote(symbol=ticker.upper())
        return quote if quote else None
    except Exception as e:
        st.error(f"Error fetching quote from Finnhub: {e}")
        return None

@st.cache_data(ttl=300)
def get_stock_metrics(ticker: str):
    """Fetches all metrics (P/E, 52 Week High/Low, etc.) from Finnhub."""
    if not finnhub_client:
        return None
    try:
        metrics = finnhub_client.company_metric(symbol=ticker.upper(), metric="all")
        return metrics if metrics else None
    except Exception as e:
        st.error(f"Error fetching metrics from Finnhub: {e}")
        return None

# ==============================
# DATA FILTERING FUNCTION (PROFESSIONAL GRADE)
# ==============================
def extract_key_financials(profile_data: dict, metric_data: dict):
    """
    Filters raw Finnhub profile and metric data to extract core institutional metrics,
    cleaning up missing values to maintain a high-end UI standard.
    """
    if not profile_data:
        profile_data = {}

    metrics = {}
    if metric_data and isinstance(metric_data, dict):
        metrics = metric_data.get("metric", {})

    # Helper to format clean numbers or default gracefully
    def safe_val(val, scale=1, suffix=""):
        if val is None or val == "N/A":
            return "N/A"
        try:
            num = float(val)
            if scale != 1:
                num = num * scale
            return f"{num:,.2f}{suffix}"
        except:
            return str(val)

    filtered_data = {
        "Company": profile_data.get("name", "N/A"),
        "Ticker": profile_data.get("ticker", "N/A"),
        "Market Cap (M)": safe_val(profile_data.get("marketCapitalization")),
        "Currency": profile_data.get("currency", "USD"),
        "P/E Ratio (TTM)": safe_val(metrics.get("peBasicExclExtraTTM")),
        "52 Week High": safe_val(metrics.get("52WeekHigh")),
        "52 Week Low": safe_val(metrics.get("52WeekLow")),
        "Dividend Yield": safe_val(metrics.get("dividendYieldIndicatedAnnual"), suffix="%"),
        "Beta": safe_val(metrics.get("beta")),
        "Profit Margin": safe_val(metrics.get("netProfitMarginTTM"), suffix="%"),
        "ROE": safe_val(metrics.get("roeTTM"), suffix="%"),
        "ROCE": safe_val(metrics.get("roceTTM"), suffix="%")
    }

    return filtered_data

# ==============================
# ADVANCED CRITICAL GEMINI AI RESEARCH FUNCTION
# ==============================
def generate_ai_research(company_name: str, ticker: str, profile_data: dict = None):
    """Generates institutional-grade, critical financial research using Gemini."""
    if not gemini_client:
        return "Error: Gemini API Key is missing. Please check your .env file."
    
    try:
        prompt = f"""
        Act as a top-tier institutional equity research analyst and senior portfolio manager. 
        Your mission is to provide an EXHAUSTIVE, deeply critical, and highly structured professional research report for: **{company_name} (Ticker: {ticker})**.
        
        USE THE FOLLOWING DATA AS CONTEXT: {profile_data}
        
        STRICT MANDATORY INSTRUCTIONS:
        1. NO SUMMARIES: Provide maximum depth, academic rigor, and professional insight. Do not truncate or shorten your analysis.
        2. TONE: Maintain a professional, authoritative, and uncompromisingly critical tone. Avoid basic textbook definitions; focus on institutional-level financial insights.
        3. FINANCIAL RIGOR: Analyze Moat, ROIC, Capital Allocation, Transfer Pricing, and Regulatory risks with high detail.
        4. CRITICALITY: Aggressively highlight vulnerabilities, structural headwinds, and macro threats.
        
        REQUIRED STRUCTURE:
        1. **Executive Summary & Investment Thesis**: Provide deep insights into the company's long-term viability, market positioning, and structural competitive advantages.
        2. **Business Model, Moat & Unit Economics**: Provide a detailed breakdown of the engine of profitability. Include a rigorous analysis of the company's economic moat.
        3. **Financial Health, Capital Allocation & Margin Dynamics**: Analyze profitability, cash flow quality, debt-to-equity leverage, and capital allocation efficiency.
        4. **Industry Headwinds & Critical Risk Factors**: Provide at least 4 deep-seated risks. Explain the quantitative and qualitative impact of each.
        5. **Growth Catalysts & Future Outlook**: Identify and analyze specific expansion vectors, innovation pipelines, and future growth drivers.
        6. **Institutional Due Diligence Checklist**: Provide a rigorous, pro-level checklist of 4-5 probing questions that a professional investor should ask the C-suite management.
        
        Ensure the output is long-form, highly detailed, and formatted using clean Markdown (headers, bullet points, and bold text for key metrics). Do not hold back on critical analysis.
        """
        
        response = gemini_client.models.generate_content(
            model="gemini-3.5-flash",
            contents=prompt,
        )
        return response.text
    
    except Exception as e:
        return f"Error generating AI analysis: {e}"

# ==============================
# WORD REPORT GENERATOR FUNCTION
# ==============================
def generate_word_report(company_name: str, ticker: str, report_text: str):
    """Generates a clean Word (.docx) file from the AI research report text."""
    doc = Document()
    
    # Title
    doc.add_heading(f"AI Investment Research Report: {company_name} ({ticker})", level=1)
    doc.add_paragraph() # Spacing
    
    # Add paragraphs from AI report cleanly
    for line in report_text.split("\n"):
        clean_line = line.strip()
        if clean_line:
            if clean_line.startswith("#") or "**" in clean_line:
                p = doc.add_paragraph()
                run = p.add_run(clean_line.replace("*", "").replace("#", ""))
                run.bold = True
            else:
                doc.add_paragraph(clean_line)
        else:
            doc.add_paragraph()
            
    # Save to BytesIO buffer so Streamlit can download it directly
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer